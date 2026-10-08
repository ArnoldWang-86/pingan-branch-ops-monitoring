# -*- coding: utf-8 -*-
"""md2docx.py —— 把 Markdown 转成中文排版正确的 Word 文档。

要点（否则中文在 Word 里会变成方框或回退字体）：
  中文必须显式设置 w:eastAsia 字体，只设 run.font.name 是不够的。

用法：
  python md2docx.py 输入.md 输出.docx
"""
from __future__ import annotations

import os
import re
import sys

# 中文 Windows 控制台默认 GBK，脚本里的非 GBK 字符（✓ ⚠ ▪ 等）
# 在 print 时会抛 UnicodeEncodeError 并中断执行。入口处重配为 UTF-8。
for _s in ("stdout", "stderr"):
    _st = getattr(sys, _s, None)
    if _st is not None and hasattr(_st, "reconfigure"):
        try:
            _st.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
PROJ = os.path.abspath(os.path.join(HERE, "..", ".."))

FONT = "微软雅黑"
MONO = "Consolas"
NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x5C, 0x9A)
GRAY = RGBColor(0x60, 0x60, 0x60)


def _tbl_width(emu):
    """生成 <w:tblW> 元素，把表格宽度固定为版心宽度。"""
    el = OxmlElement("w:tblW")
    el.set(qn("w:w"), str(int(emu / 635)))     # EMU -> twips
    el.set(qn("w:type"), "dxa")
    return el


def _tbl_cell_margins(table, left=0, right=0, top=20, bottom=20):
    """显式设定单元格内边距。

    必须设：Word/LibreOffice 的表格默认左右各留约 108 twips 内边距，
    它是在列宽之外额外占位的。若列宽之和已等于版心宽度，
    这些内边距就会把整张表推出右边距，导致右侧文字被裁切
    （这是渲染出图后才发现的问题，只看 XML 看不出来）。
    """
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    table._tbl.tblPr.append(mar)


def set_cjk(run, font=FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), font)
    return run


def emit_runs(paragraph, text, *, size=None, italic=False, color=None, force_bold=False):
    """写入一段文本，处理 **粗体** 与 `行内代码` 两种标记。"""
    for seg in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not seg:
            continue
        bold = force_bold

        if seg.startswith("**") and seg.endswith("**"):
            seg, bold = seg[2:-2], True
            run = paragraph.add_run(seg)
        elif seg.startswith("`") and seg.endswith("`"):
            run = paragraph.add_run(seg[1:-1])
            run.font.name = MONO
            rpr = run._element.get_or_add_rPr()
            rf = rpr.find(qn("w:rFonts"))
            if rf is None:
                rf = OxmlElement("w:rFonts")
                rpr.insert(0, rf)
            for attr in ("w:ascii", "w:hAnsi"):
                rf.set(qn(attr), MONO)
            rf.set(qn("w:eastAsia"), FONT)
            run.font.size = Pt((size or 10.5) - 1)
            if color is not None:
                run.font.color.rgb = color
            if italic:
                run.italic = True
            continue
        else:
            run = paragraph.add_run(seg)

        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if italic:
            run.italic = True
        if color is not None:
            run.font.color.rgb = color
        set_cjk(run)
    return paragraph


def convert(src: str, dst: str) -> tuple[int, int]:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.28

    # 版心宽度：表格需要按它来限制列宽，否则会撑出右边距
    from docx.shared import Mm
    sec = doc.sections[0]
    sec_width = sec.page_width
    sec_left = sec.left_margin
    sec_right = sec.right_margin

    def heading(text, level):
        sizes = {0: 18, 1: 15, 2: 13, 3: 11.5}
        colors = {0: NAVY, 1: NAVY, 2: BLUE, 3: GRAY}
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14 if level <= 1 else 9)
        p.paragraph_format.space_after = Pt(6)
        emit_runs(p, text, size=sizes.get(level, 11),
                  color=colors.get(level, RGBColor(0, 0, 0)), force_bold=True)
        return p

    def code_block(lines):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(8)
        r = p.add_run("\n".join(lines))
        r.font.name = MONO
        r.font.size = Pt(9)
        r.font.color.rgb = NAVY
        rpr = r._element.get_or_add_rPr()
        rf = rpr.find(qn("w:rFonts"))
        if rf is not None:
            rf.set(qn("w:eastAsia"), FONT)
        return p

    def table(rows):
        ncol = max(len(r) for r in rows)
        t = doc.add_table(rows=0, cols=ncol)
        try:
            t.style = "Light Grid Accent 1"
        except KeyError:
            t.style = "Table Grid"
        # 表格必须显式限制在版心宽度内，否则 LibreOffice/Word 会按内容撑开，
        # 超出右边距导致文字被裁切（这是实际渲染后才发现的问题）。
        t.autofit = False
        usable = sec_width - sec_left - sec_right
        col_w = int(usable / ncol)
        # 先固定表格宽度与单元格内边距（内边距会额外占位，必须设小），
        # 再把列宽之和控制在版心宽度内。
        t._tbl.tblPr.append(_tbl_width(usable))
        _tbl_cell_margins(t, left=60, right=60)
        for ci, col in enumerate(t.columns):
            col.width = Emu(col_w)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            for ci in range(ncol):
                val = row[ci] if ci < len(row) else ""
                cells[ci].width = Emu(col_w)
                cells[ci].text = ""
                p = cells[ci].paragraphs[0]
                p.paragraph_format.space_after = Pt(2)
                # 表格内文字略小，并允许中文换行
                emit_runs(p, val, size=9, force_bold=(ri == 0))
        doc.add_paragraph()
        return t

    def quote(text):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        emit_runs(p, text, size=10, italic=True, color=GRAY)
        return p

    def bullet(text, level=0):
        try:
            p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        except KeyError:
            p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        emit_runs(p, text)
        return p

    def numbered(text):
        try:
            p = doc.add_paragraph(style="List Number")
        except KeyError:
            p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        emit_runs(p, text)
        return p

    lines = open(src, encoding="utf-8").read().split("\n")
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        if line.startswith("```"):
            buf, i = [], i + 1
            while i < len(lines) and not lines[i].startswith("```"):
                buf.append(lines[i])
                i += 1
            code_block(buf)
            i += 1
            continue

        if line.startswith("|") and i + 1 < len(lines) \
                and re.match(r"^\|[\s\-:|]+\|", lines[i + 1]):
            rows = [[c.strip() for c in line.strip("|").split("|")]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table(rows)
            continue

        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            heading(m.group(2).strip(), len(m.group(1)) - 1)
            i += 1
            continue

        if line.startswith(">"):
            quote(line.lstrip("> ").strip())
            i += 1
            continue

        if re.match(r"^-{3,}$", line.strip()):
            p = doc.add_paragraph()
            r = p.add_run("─" * 40)
            r.font.color.rgb = RGBColor(0xC0, 0xC0, 0xC0)
            set_cjk(r)
            i += 1
            continue

        m = re.match(r"^(\s*)[-*]\s+(.*)", line)
        if m:
            bullet(m.group(2), 1 if len(m.group(1)) >= 2 else 0)
            i += 1
            continue

        m = re.match(r"^(\s*)(\d+)\.\s+(.*)", line)
        if m:
            numbered(m.group(3))
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        p = doc.add_paragraph()
        emit_runs(p, line)
        i += 1

    doc.save(dst)
    return len(doc.paragraphs), len(doc.tables)


def build_all():
    """按目录自动发现 Markdown 并转成同名 Word。

    为什么要提供这个入口：在中文 Windows 下，如果从 .bat 里把中文文件名
    作为命令行参数传给 Python，参数会经过 cmd.exe 的代码页转换而损坏
    （实测报 FileNotFoundError，文件名里全是 \\ufffd）。让脚本自己在
    目录里找文件，就完全绕开了这一层编码问题——目录名是 ASCII 的。
    """
    pairs = [
        (os.path.join(PROJ, "results"), "运营建议报告.md"),
        (os.path.join(PROJ, "docs"), "方法说明书.md"),
        (PROJ, "README.md"),
    ]
    ok = 0
    for folder, name in pairs:
        src = os.path.join(folder, name)
        if not os.path.exists(src):
            print(f"跳过（未找到）: {src}")
            continue
        dst = os.path.splitext(src)[0] + ".docx"
        np_, nt = convert(src, dst)
        print(f"已生成: {dst}（段落 {np_} / 表格 {nt}）")
        ok += 1
    print(f"\n共转换 {ok} 份文档。")


def main():
    if len(sys.argv) >= 2 and sys.argv[1] == "--all":
        build_all()
        return
    if len(sys.argv) < 3:
        print(__doc__)
        print("也可运行： python md2docx.py --all   （自动转换 results/docs/README）")
        raise SystemExit(1)
    src, dst = sys.argv[1], sys.argv[2]
    np_, nt = convert(src, dst)
    print(f"已生成: {dst}（段落 {np_} / 表格 {nt}）")


if __name__ == "__main__":
    main()
